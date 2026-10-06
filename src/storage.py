"""
Data storage module

Responsibilities:
  - Persist the collected results to local files
  - Support two formats: JSON (complete raw data) and Excel (multi-sheet workbook)
  - Name files by keyword and timestamp to avoid overwriting

Directory layout:
    data/
    ├── raw/
    │   ├── {keyword}_{timestamp}.json         # Raw search result data
    │   └── notes_{keyword}_{timestamp}.json   # Raw note detail data
    └── processed/
        └── {keyword}_{timestamp}.xlsx         # Excel summary (3 sheets)

Usage:
    storage = Storage(config["storage"])
    storage.save_all("coffee", search_results, note_details)
"""

from __future__ import annotations

import json
import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# Column header definitions for each Excel sheet
_SEARCH_FIELDS = [
    "note_id",
    "title",
    "author",
    "author_id",
    "likes",
    "note_type",
    "note_url",
    "publish_time",
]

_NOTE_FIELDS = [
    "note_id",
    "title",
    "content",
    "author",
    "author_id",
    "publish_time",
    "likes",
    "collects",
    "comments_count",
    "shares",
    "tags",
    "note_type",
    "note_url",
]

_COMMENT_FIELDS = [
    "comment_id",
    "note_id",
    "user_name",
    "user_id",
    "content",
    "likes",
    "time",
    "ip_location",
]

# Display header for each field key
_HEADER_LABELS = {
    "note_id": "Note ID",
    "title": "Title",
    "content": "Content",
    "author": "Author",
    "author_id": "Author ID",
    "publish_time": "Publish Time",
    "likes": "Likes",
    "collects": "Collects",
    "comments_count": "Comment Count",
    "shares": "Shares",
    "tags": "Tags",
    "note_type": "Note Type",
    "note_url": "Note URL",
    "comment_id": "Comment ID",
    "user_name": "User",
    "user_id": "User ID",
    "time": "Time",
    "ip_location": "IP Location",
}


class Storage:
    """Local data storage manager.

    Decides from the config whether to write JSON / Excel, and handles directory creation and file naming.
    """

    def __init__(self, config: dict) -> None:
        """Initialise the storage config.

        Args:
            config: Dict of the storage node in settings.yaml, containing:
                - output_dir (str): Output root directory, default "data"
                - save_raw_json (bool): Whether to save the raw JSON
                - save_xlsx (bool): Whether to save the Excel file
        """
        self._root = Path(config.get("output_dir", "data"))
        self._save_json: bool = config.get("save_raw_json", True)
        self._save_xlsx: bool = config.get("save_xlsx", True)
        self._ensure_dirs()

    def _ensure_dirs(self) -> None:
        """Make sure the output directories exist."""
        (self._root / "raw").mkdir(parents=True, exist_ok=True)
        (self._root / "processed").mkdir(parents=True, exist_ok=True)

    def save_all(
        self,
        keyword: str,
        search_results: list[dict],
        note_details: list[dict],
    ) -> None:
        """Save all collected data in one go (JSON + Excel).

        Args:
            keyword: Search keyword (used for file naming)
            search_results: List of dicts returned by parse_search_card()
            note_details: List of note details returned by fetch_note_details(),
                          each containing the detail fields + a comments sub-list
        """
        safe_keyword = _sanitize_filename(keyword)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        # Write JSON
        if self._save_json:
            if search_results:
                self._write_json(safe_keyword, timestamp, keyword, search_results)
            if note_details:
                self._write_notes_json(safe_keyword, timestamp, keyword, note_details)

        # Write Excel
        if self._save_xlsx:
            self._write_xlsx(safe_keyword, timestamp, search_results, note_details)

    # ---- JSON writers ----

    def _write_json(
        self,
        safe_keyword: str,
        timestamp: str,
        keyword: str,
        results: list[dict],
    ) -> None:
        """Write the search results JSON file."""
        json_path = self._root / "raw" / f"{safe_keyword}_{timestamp}.json"
        payload = {
            "keyword": keyword,
            "crawled_at": datetime.now().isoformat(timespec="seconds"),
            "count": len(results),
            "results": results,
        }
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        logger.info("JSON written: %s (%d items)", json_path, len(results))

    def _write_notes_json(
        self,
        safe_keyword: str,
        timestamp: str,
        keyword: str,
        note_details: list[dict],
    ) -> None:
        """Write the note details JSON file (including comments)."""
        json_path = self._root / "raw" / f"notes_{safe_keyword}_{timestamp}.json"
        payload = {
            "keyword": keyword,
            "crawled_at": datetime.now().isoformat(timespec="seconds"),
            "count": len(note_details),
            "notes": note_details,
        }
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        logger.info("Note details JSON written: %s (%d items)", json_path, len(note_details))

    # ---- Excel writers ----

    def _write_xlsx(
        self,
        safe_keyword: str,
        timestamp: str,
        search_results: list[dict],
        note_details: list[dict],
    ) -> None:
        """Generate an Excel file with 3 sheets.

        Sheet layout:
          - Search Results: note summaries obtained during the search phase
          - Note Details: note body, engagement data, etc.
          - Comments: comments from all notes combined
        """
        wb = Workbook()

        # Sheet 1: Search Results
        ws_search = wb.active
        ws_search.title = "Search Results"
        self._fill_sheet(ws_search, _SEARCH_FIELDS, search_results)

        # Sheet 2: Note Details (tags list joined into a string, nested fields removed)
        ws_notes = wb.create_sheet("Note Details")
        note_rows = []
        for note in note_details:
            row = dict(note)
            row["tags"] = ";".join(row.get("tags", []))
            row.pop("comments", None)
            row.pop("images", None)
            row.pop("video_url", None)
            note_rows.append(row)
        self._fill_sheet(ws_notes, _NOTE_FIELDS, note_rows)

        # Sheet 3: Comments combined
        ws_comments = wb.create_sheet("Comments")
        all_comments: list[dict] = []
        for note in note_details:
            all_comments.extend(note.get("comments", []))
        self._fill_sheet(ws_comments, _COMMENT_FIELDS, all_comments)

        # Save the file
        xlsx_path = self._root / "processed" / f"{safe_keyword}_{timestamp}.xlsx"
        wb.save(xlsx_path)
        logger.info(
            "Excel written: %s (%d search results / %d notes / %d comments)",
            xlsx_path,
            len(search_results),
            len(note_details),
            len(all_comments),
        )

    def _fill_sheet(
        self,
        ws,
        fieldnames: list[str],
        rows: list[dict],
    ) -> None:
        """Fill a single sheet: write the header + data rows + formatting.

        Formatting covers: frozen first row, auto-filter, auto-fitted column widths.
        """
        # Write the header
        headers = [_HEADER_LABELS.get(field, field) for field in fieldnames]
        ws.append(headers)

        # Write the data rows
        for row in rows:
            ws.append([row.get(field) for field in fieldnames])

        # Freeze the first row (header stays visible while scrolling)
        ws.freeze_panes = "A2"

        # Auto-filter (covers all data columns)
        if rows:
            last_col = get_column_letter(len(fieldnames))
            last_row = len(rows) + 1  # +1 for the header row
            ws.auto_filter.ref = f"A1:{last_col}{last_row}"

        # Auto-fit column widths (based on the longest of the header and the content)
        for col_idx, field in enumerate(fieldnames, start=1):
            # Compute the column's maximum character width (header + a sample of the first 100 data rows)
            max_len = len(headers[col_idx - 1])
            for row in rows[:100]:
                val = row.get(field)
                if val is not None:
                    # Non-ASCII (e.g. CJK) characters count as double width
                    cell_len = sum(2 if ord(c) > 127 else 1 for c in str(val))
                    max_len = max(max_len, cell_len)
            # Cap the column width at 60, with a minimum of 10
            col_width = min(max(max_len + 2, 10), 60)
            ws.column_dimensions[get_column_letter(col_idx)].width = col_width


def _sanitize_filename(name: str) -> str:
    """Turn a string into a safe filename (strips special characters such as / \\ : * ? " < > |).

    Args:
        name: The original string

    Returns:
        A safe filename string (CJK characters, letters, digits, underscores and hyphens are kept)
    """
    import re
    # Replace unsafe characters with underscores
    safe = re.sub(r'[\\/:*?"<>|\s]', "_", name)
    # Collapse consecutive underscores
    safe = re.sub(r"_+", "_", safe)
    return safe.strip("_") or "unnamed"
